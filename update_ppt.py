import pptx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

PPT_PATH = "ADS Project PPT.pptx"

if not os.path.exists(PPT_PATH):
    print(f"Error: {PPT_PATH} not found!")
    sys.exit(1)

prs = pptx.Presentation(PPT_PATH)

def set_slide_text(slide, shape_texts):
    for shape_idx, text_list in shape_texts.items():
        if shape_idx < len(slide.shapes):
            shape = slide.shapes[shape_idx]
            if shape.has_text_frame:
                tf = shape.text_frame
                tf.clear()
                for p_idx, line in enumerate(text_list):
                    if p_idx == 0:
                        p = tf.paragraphs[0]
                    else:
                        p = tf.add_paragraph()
                    p.text = line

# SLIDE 1: Title & Presentation
set_slide_text(prs.slides[0], {
    0: [
        "Sanjivani Rural Education Society's",
        "Sanjivani College of Engineering, Kopargaon-423603",
        "Department of Computer Engineering"
    ],
    3: [
        "A Presentation On Mini Project - Applied Data Science Laboratory (PCCO314A)",
        "\"Household Electricity Consumption Prediction Using ANN\"",
        "BE / B.Tech 2025-26 SEM VII",
        "",
        "Presented By:",
        "1. Mundhe Aditya",
        "2. Borude Yash Ambadas",
        "3. Langhi Roshan",
        "4. Shinde Tushar Maruti",
        "5. Chaudhari Abhishek Anil"
    ],
    4: [
        "Guided By:",
        "Prof. S. A. Shivarkar Sir",
        "Department of Computer Engineering"
    ]
})

# SLIDE 2: Problem Definition (Mundhe Aditya)
set_slide_text(prs.slides[1], {
    1: ["DEPARTMENT OF COMPUTER ENGINEERING, Sanjivani COE, Kopargaon"],
    2: ["Problem Definition"],
    3: [
        "Energy Grid Dynamics: Rapid fluctuations in household power consumption cause grid load imbalances.",
        "Peak Demand Penalties: Unpredicted consumption spikes result in higher utility costs and equipment stress.",
        "High Telemetry Volume: Raw smart meter data (2.07M+ 1-minute records) requires automated real-time processing.",
        "Inefficient Static Models: Traditional linear models fail to capture complex non-linear electrical relationships."
    ]
})

# SLIDE 3: Objectives (Mundhe Aditya)
set_slide_text(prs.slides[2], {
    0: [
        "Telemetry Data Ingestion: Process 2.07M+ minute-level records from UCI Household Power dataset.",
        "Feature Engineering & Scaling: Extract temporal features (Hour, Day, Month, DayOfWeek) & scale inputs.",
        "Deep Neural Network Training: Train a Deep ANN with Batch Normalization and Dropout (0.2).",
        "High Prediction Accuracy: Achieve high accuracy (R² > 99%) for instantaneous power consumption (kW).",
        "Interactive Web Deployment: Build a theme-adaptive Streamlit Dashboard with tariff cost calculation & batch prediction."
    ],
    1: ["DEPARTMENT OF COMPUTER ENGINEERING, Sanjivani COE, Kopargaon"],
    2: ["Objectives"]
})

# SLIDE 4: Dataset Description (Borude Yash Ambadas)
set_slide_text(prs.slides[3], {
    1: ["DEPARTMENT OF COMPUTER ENGINEERING, Sanjivani COE, Kopargaon"],
    2: ["Dataset Description"],
    4: [
        "Target Variable:",
        "Global_active_power (kW) - Household active electrical power"
    ],
    5: [
        "Key Input Features:",
        "Global_reactive_power (kVAR) - Household reactive power",
        "Voltage (V) - Measured household voltage",
        "Global_intensity (A) - Current intensity in amperes",
        "Sub_metering_1 (Wh) - Kitchen appliances (dishwasher, oven)",
        "Sub_metering_2 (Wh) - Laundry room (washing machine, fridge)",
        "Sub_metering_3 (Wh) - Climate control (AC, water heater)"
    ]
})

# SLIDE 5: Feature Engineering & Preprocessing (Borude Yash Ambadas)
set_slide_text(prs.slides[4], {
    0: ["Feature Engineering & Preprocessing"],
    1: [
        "1. Temporal Features Extracted",
        "Hour of Day (0-23)",
        "Day of Month (1-31)",
        "Month (1-12)",
        "Day of Week (Monday-Sunday)"
    ],
    2: [
        "2. Data Cleaning",
        "Converted missing '?' values to NaN",
        "Numeric datatype casting",
        "Chronological sorting by Datetime"
    ],
    3: [
        "3. Scaling & Train Split",
        "Applied StandardScaler to prevent leakage",
        "80% Training / 20% Testing split"
    ],
    4: [
        "Pipeline Output: Cleaned, scaled feature matrix ready for neural network training."
    ]
})

# SLIDE 6: Model Development (Langhi Roshan)
set_slide_text(prs.slides[5], {
    0: ["Model Development (ANN Architecture)"],
    1: [
        "Deep Neural Network Architecture:",
        "Input Layer: 10 Scaled Features",
        "Dense Layer 1: 64 Neurons (ReLU activation)",
        "Batch Normalization + Dropout (0.2 rate)",
        "Dense Layer 2: 32 Neurons (ReLU activation) + Dropout (0.2)",
        "Dense Layer 3: 16 Neurons (ReLU activation)",
        "Output Layer: 1 Neuron (Linear activation for kW regression)",
        "",
        "Optimizer & Loss Function: Adam Optimizer with Mean Squared Error (MSE).",
        "Model Saving: Saved as electricity_ann.keras and lightweight ann_weights.pkl."
    ]
})

# SLIDE 7: System Architecture (Langhi Roshan)
set_slide_text(prs.slides[6], {
    1: ["DEPARTMENT OF COMPUTER ENGINEERING, Sanjivani COE, Kopargaon"],
    2: ["System Architecture"],
    3: ["End-to-End System Pipeline & Workflow"],
    5: ["Raw Telemetry -> Preprocessing & Scaling -> Deep ANN Engine -> Streamlit Web UI & Tariff Calculator"]
})

# SLIDE 8: Result Analysis (Shinde Tushar Maruti)
set_slide_text(prs.slides[7], {
    1: ["DEPARTMENT OF COMPUTER ENGINEERING, Sanjivani COE, Kopargaon"],
    2: ["Result Analysis & Metrics"],
    4: [
        "Quantitative Evaluation Metrics:",
        "Mean Absolute Error (MAE): 0.0573 kW (~57 Watts error)",
        "Mean Squared Error (MSE): 0.0061 kW²",
        "Root Mean Squared Error (RMSE): 0.0783 kW",
        "R² Accuracy Score: 0.9943 (99.43% variance explained)"
    ],
    5: [
        "Key Finding: Deep ANN with Batch Normalization achieves exceptional precision in tracking rapid household electrical load changes."
    ]
})

# SLIDE 9: Conclusion (Shinde Tushar Maruti)
set_slide_text(prs.slides[8], {
    1: ["DEPARTMENT OF COMPUTER ENGINEERING, Sanjivani COE, Kopargaon"],
    2: ["Conclusion"],
    3: [
        "High Forecasting Accuracy: ANN model achieves R² = 0.9943 with MAE of 0.0573 kW.",
        "Effective Regularization: Batch Normalization and Dropout prevented overfitting across 2M+ samples.",
        "Real-Time Analytics: Streamlit app provides real-time single & batch CSV predictions.",
        "Tariff Estimation: Integrated Indian electricity cost calculator (₹/kWh) for consumer utility tracking."
    ]
})

# SLIDE 10: References & Testing (Chaudhari Abhishek Anil)
set_slide_text(prs.slides[9], {
    1: ["DEPARTMENT OF COMPUTER ENGINEERING, Sanjivani COE, Kopargaon"],
    2: ["References & Quality Assurance"],
    3: [
        "Dataset Source: UCI Machine Learning Repository (Individual Household Electric Power Consumption).",
        "Machine Learning Framework: TensorFlow / Keras & Scikit-Learn Documentation.",
        "Web Application & Visualization: Streamlit & Plotly Express.",
        "Automated Testing: Built-in unittest suite (test_model.py) verifying scaler and ANN forward pass."
    ]
})

# SLIDE 11: Thank You
set_slide_text(prs.slides[10], {
    0: ["THANK YOU!!", "Questions & Discussion"]
})

prs.save(PPT_PATH)
print("Successfully updated PPT content in ADS Project PPT.pptx!")
